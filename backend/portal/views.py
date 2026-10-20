import base64
import json
import secrets
import uuid
from io import BytesIO
from datetime import date, datetime, timedelta
from django.contrib.auth import authenticate, login, logout, password_validation
from django.db import transaction, IntegrityError
from django.db.models import Count, Q
from django.http import FileResponse, Http404
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.csrf import ensure_csrf_cookie, csrf_protect
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.paginator import Paginator
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from .models import (
    User, Department, Lecturer, DepartmentHistory, EmploymentRecord, ScienceProfile,
    EducationRecord, WorkType, ResearchWork, WorkAuthor, WorkEvidence, WorkReview,
    WorkLinkRequest, WorkCorrection, WorkVersion, IntellectualPropertyRecord,
    PracticalApplicationRecord, ResearchTaskRecord, AwardRecord,
    ProfileApproval, ApprovalEvent, AuditLog, Notification, SigningIdentity, DigitalSignature,
    EvidenceDocument, EvidenceLink, VerificationRecord, ApprovalVerificationItem,
)
from .services import (
    normalize_title, normalize_doi, find_duplicate, store_private, store_bytes,
    read_private, snapshot_for, snapshot_hash, audit, stable_hash, evidence_for, evidence_hash, evidence_payload, work_evidence_for,
)
from .pdf_service import profile_pdf
from .signing_service import provision_identity, identity_public_payload, sign_bytes, verify_bytes, sha256_bytes
from .verification_service import create_approval_verification_items, review_item, all_items_verified, evidence_usable


def failure(message, status_code=400, **extra):
    return Response({'error': message, **extra}, status=status_code)


def has_role(user, *roles):
    return user.is_authenticated and (user.role in roles or (user.is_superuser and 'admin' in roles))


def require_role(request, *roles):
    # Admin is the system-wide operator in V1 and may execute every protected workflow.
    if request.user.is_authenticated and (request.user.role == 'admin' or request.user.is_superuser):
        return None
    return None if has_role(request.user, *roles) else failure('Không có quyền thực hiện thao tác.', 403)


def own_lecturer(request):
    return Lecturer.objects.select_related('department', 'user').filter(user=request.user).first()

def target_lecturer(request, data=None):
    """Resolve the lecturer operated on. Admin/viewer_admin may select any lecturer; mutation is guarded separately."""
    if request.user.role in ('admin', 'viewer_admin') or request.user.is_superuser:
        source = data or {}
        lecturer_id = source.get('lecturer_id') or request.query_params.get('lecturer_id')
        if not lecturer_id:
            return None
        return Lecturer.objects.select_related('department', 'user').filter(pk=lecturer_id).first()
    return own_lecturer(request)


def json_body(request):
    return request.data if isinstance(request.data, dict) else {}


def error_validation(exc):
    return failure(str(exc.messages[0] if hasattr(exc, 'messages') else exc))


def date_value(value, required=False):
    if not value:
        if required:
            raise ValueError('Thiếu ngày bắt buộc.')
        return None
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def year_value(value, required=False):
    if value in (None, ''):
        if required:
            raise ValueError('Thiếu năm bắt buộc.')
        return None
    year = int(value)
    if not 1900 <= year <= timezone.localdate().year + 2:
        raise ValueError('Năm không hợp lệ.')
    return year


def lecturer_out(x):
    return {
        'id': str(x.id), 'staff_code': x.staff_code, 'full_name': x.full_name,
        'personal_id': x.personal_id, 'date_of_birth': str(x.date_of_birth) if x.date_of_birth else '',
        'gender': x.gender, 'gender_label': x.get_gender_display() if x.gender else '',
        'professional_title': x.professional_title, 'position': x.position,
        'home_address': x.home_address, 'phone': x.phone, 'email': x.email,
        'department_id': x.department_id, 'department': x.department.name,
        'status': x.status, 'username': x.user.username,
    }


def work_out(w):
    return {
        'id': str(w.id), 'title': w.title, 'work_type_id': w.work_type_id,
        'work_type': w.work_type.name, 'publication_year': w.publication_year,
        'doi': w.doi, 'venue': w.venue, 'status': w.status,
        'created_by_id': w.created_by_id, 'created_at': w.created_at.isoformat(),
        'authors': [{
            'id': a.pk, 'name': a.name,
            'lecturer_id': str(a.lecturer_id) if a.lecturer_id else None,
            'author_order': a.author_order, 'authorship_type': a.authorship_type,
            'authorship_label': a.get_authorship_type_display(), 'role_note': a.role_note,
        } for a in w.authors.all()],
        # Từ V6.4 đây là minh chứng công trình khoa học bắt buộc, không còn là
        # tài liệu phụ trợ tùy chọn.
        'evidence': work_evidence_for(w),
    }


def approval_out(a):
    total_items = a.verification_items.count() if hasattr(a, 'verification_items') else 0
    verified_items = a.verification_items.filter(status=ApprovalVerificationItem.Status.VERIFIED).count() if hasattr(a, 'verification_items') else 0
    return {
        'id': str(a.id), 'code': a.code, 'version': a.version,
        'parent_approval_id': str(a.parent_approval_id) if a.parent_approval_id else None,
        'form_code': a.form_code, 'form_version': a.form_version, 'lecturer_id': str(a.lecturer_id),
        'lecturer': a.lecturer.full_name, 'task_name': a.task_name,
        'registered_role': a.registered_role, 'registered_role_label': a.get_registered_role_display(),
        'registered_role_other': a.registered_role_other,
        'status': a.status, 'comment': a.comment, 'submitted_at': a.submitted_at.isoformat(),
        'has_scan': bool(a.scan_path), 'has_signed_pdf': bool(a.signed_pdf_path),
        'has_personal_signature': a.digital_signatures.filter(signature_role=DigitalSignature.SignatureRole.PERSONAL).exists(),
        'has_unit_signature': a.digital_signatures.filter(signature_role=DigitalSignature.SignatureRole.UNIT).exists(),
        'verification_total': total_items, 'verification_verified': verified_items,
        'verification_complete': total_items == verified_items,
        'approved_at': a.approved_at.isoformat() if a.approved_at else None,
        'completed_at': a.completed_at.isoformat() if a.completed_at else None,
    }


def can_access_work(request, w):
    if has_role(request.user, 'admin', 'viewer_admin', 'reviewer', 'signer'):
        return True
    return w.status == ResearchWork.Status.VERIFIED or w.created_by_id == request.user.id or w.authors.filter(lecturer__user=request.user).exists()


@ensure_csrf_cookie
@api_view(['GET'])
@permission_classes([AllowAny])
def csrf(request):
    return Response({'ok': True})


@csrf_protect
@api_view(['POST'])
@permission_classes([AllowAny])
def sign_in(request):
    d = json_body(request)
    user = authenticate(request, username=d.get('username', ''), password=d.get('password', ''))
    if not user:
        return failure('Tên đăng nhập hoặc mật khẩu không đúng.', 401)
    login(request, user)
    audit(user, 'login', 'user', user.id)
    return Response(user_payload(user))


@api_view(['POST'])
def sign_out(request):
    audit(request.user, 'logout', 'user', request.user.id)
    logout(request)
    return Response({'ok': True})


def user_payload(user):
    lecturer = Lecturer.objects.select_related('department', 'user').filter(user=user).first()
    return {
        'id': user.id, 'username': user.username, 'role': user.role,
        'must_change_password': user.must_change_password,
        'lecturer': lecturer_out(lecturer) if lecturer else None,
    }


@api_view(['GET'])
def me(request):
    return Response(user_payload(request.user))


@api_view(['POST'])
def change_password(request):
    d = json_body(request)
    if not request.user.check_password(d.get('current_password', '')):
        return failure('Mật khẩu hiện tại không đúng.')
    new = d.get('new_password', '')
    try:
        password_validation.validate_password(new, request.user)
    except DjangoValidationError as ex:
        return error_validation(ex)
    request.user.set_password(new)
    request.user.must_change_password = False
    request.user.save(update_fields=['password', 'must_change_password'])
    from django.contrib.auth import update_session_auth_hash
    update_session_auth_hash(request, request.user)
    audit(request.user, 'password_changed', 'user', request.user.id)
    return Response({'ok': True})


@api_view(['GET'])
def catalogs(request):
    return Response({
        'departments': list(Department.objects.values('id', 'code', 'name', 'address', 'head_name', 'phone').order_by('name')),
        'work_types': list(WorkType.objects.filter(is_active=True).values('id', 'code', 'name').order_by('id')),
        'talent_categories': [{'value': x, 'label': y} for x, y in ScienceProfile.TalentCategory.choices],
        'registered_roles': [{'value': x, 'label': y} for x, y in ProfileApproval.RegisteredRole.choices],
    })


def _profile_evidence(l, section, entity_id=''):
    # evidence_for() đã trả đầy đủ metadata hiệu lực của minh chứng
    # (validity_mode, valid_from, valid_until, validity_status, days_remaining).
    # Không được rút gọn payload ở đây, nếu không frontend sẽ hiểu thiếu validity_mode
    # là minh chứng vô thời hạn.
    return evidence_for(l, section, entity_id)


def _work_for_profile(l, auth):
    w = work_out(auth.work) | {
        'my_authorship_type': auth.authorship_type,
        'my_authorship_label': auth.get_authorship_type_display(),
    }
    # Mục 10 dùng trực tiếp WorkEvidence làm minh chứng xác minh. Nhờ vậy file
    # tải tại trang Công trình khoa học xuất hiện đúng trong BM-04/NCKH.
    w['verification_evidence'] = w.get('evidence', [])
    return w


def profile_payload(l):
    s, _ = ScienceProfile.objects.get_or_create(lecturer=l)
    return {
        'lecturer': lecturer_out(l),
        'organization': {
            'id': l.department_id, 'name': l.department.name, 'address': l.department.address,
            'head_name': l.department.head_name, 'phone': l.department.phone,
        },
        'science': {
            'academic_rank': s.academic_rank, 'academic_rank_year': s.academic_rank_year,
            'academic_degree': s.academic_degree, 'academic_degree_year': s.academic_degree_year,
            'talent_category': s.talent_category, 'other_achievements': s.other_achievements,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.ACHIEVEMENT, 'science'),
        },
        'education': [{
            'id': x.pk, 'level': x.level, 'institution': x.institution,
            'specialization': x.specialization, 'graduation_year': x.graduation_year,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.EDUCATION, x.pk),
        } for x in l.education_records.all()],
        'employment': [{
            'id': x.pk, 'start_year': x.start_year, 'end_year': x.end_year,
            'position': x.position, 'organization': x.organization,
            'organization_address': x.organization_address,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.EMPLOYMENT, x.pk),
        } for x in l.employment_records.all()],
        'ip_records': [{
            'id': x.pk, 'title_content': x.title_content, 'grant_year': x.grant_year,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.IP, x.pk),
        } for x in l.ip_records.all()],
        'practical_applications': [{
            'id': x.pk, 'work_name': x.work_name, 'application_detail': x.application_detail,
            'start_date': str(x.start_date) if x.start_date else '', 'end_date': str(x.end_date) if x.end_date else '',
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.PRACTICAL, x.pk),
        } for x in l.practical_applications.all()],
        'research_tasks': [{
            'id': x.pk, 'name': x.name, 'participation_role': x.participation_role,
            'start_date': str(x.start_date) if x.start_date else '', 'end_date': str(x.end_date) if x.end_date else '',
            'program_name': x.program_name, 'status': x.status,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.TASK, x.pk),
        } for x in l.research_tasks.all()],
        'awards': [{
            'id': x.pk, 'description': x.description, 'award_year': x.award_year,
            'verification_evidence': _profile_evidence(l, EvidenceDocument.Section.AWARD, x.pk),
        } for x in l.awards.all()],
        'my_works': [_work_for_profile(l, a) for a in WorkAuthor.objects.select_related('work__work_type').prefetch_related('work__authors', 'work__evidence').filter(lecturer=l)],
        'verified_works': [_work_for_profile(l, a) for a in WorkAuthor.objects.select_related('work__work_type').prefetch_related('work__authors', 'work__evidence').filter(lecturer=l, work__status='verified')],
    }


def _sync_science_profile(l, d):
    s, _ = ScienceProfile.objects.get_or_create(lecturer=l)
    s.academic_rank = str(d.get('academic_rank', ''))[:150]
    s.academic_degree = str(d.get('academic_degree', ''))[:150]
    s.academic_rank_year = year_value(d.get('academic_rank_year'))
    s.academic_degree_year = year_value(d.get('academic_degree_year'))
    talent = d.get('talent_category', ScienceProfile.TalentCategory.NONE)
    if talent not in ScienceProfile.TalentCategory.values:
        raise ValueError('Nhóm nhân tài không hợp lệ.')
    s.talent_category = talent
    s.other_achievements = str(d.get('other_achievements', ''))[:20000]
    s.save()


def _sync_profile_lists(l, payload):
    """Đồng bộ danh sách nhưng giữ nguyên PK của dòng cũ.

    V5 xóa rồi tạo lại toàn bộ dòng mỗi lần lưu, khiến ID thay đổi và kết quả xác minh
    không thể tái sử dụng. V6 cập nhật theo id; chỉ xóa dòng thật sự bị người dùng bỏ.
    Minh chứng generic không bị xóa theo để BM-04 cũ vẫn truy xuất được.
    """
    education = payload.get('education', [])
    employment = payload.get('employment', [])
    ip_records = payload.get('ip_records', [])
    practical = payload.get('practical_applications', [])
    tasks = payload.get('research_tasks', [])
    awards = payload.get('awards', [])
    for name, data in [('education', education), ('employment', employment), ('ip_records', ip_records), ('practical_applications', practical), ('research_tasks', tasks), ('awards', awards)]:
        if not isinstance(data, list) or len(data) > 200:
            raise ValueError(f'Danh sách {name} không hợp lệ hoặc quá dài.')

    def finish(manager, kept):
        if kept:
            manager.exclude(pk__in=kept).delete()
        else:
            manager.all().delete()

    kept=[]
    for i,x in enumerate(education):
        level=str(x.get('level','')).strip(); institution=str(x.get('institution','')).strip(); specialization=str(x.get('specialization','')).strip()
        if not (level or institution or specialization): continue
        obj=l.education_records.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or EducationRecord(lecturer=l)
        obj.level=level[:150]; obj.institution=institution[:300]; obj.specialization=specialization[:250]
        obj.graduation_year=year_value(x.get('graduation_year'), required=True); obj.sort_order=i; obj.save(); kept.append(obj.pk)
    finish(l.education_records, kept)

    kept=[]
    for i,x in enumerate(employment):
        position=str(x.get('position','')).strip(); organization=str(x.get('organization','')).strip(); organization_address=str(x.get('organization_address','')).strip()
        if not (position or organization or organization_address or x.get('start_year') or x.get('end_year')): continue
        start_year=year_value(x.get('start_year'), required=True); end_year=year_value(x.get('end_year'))
        if end_year and end_year < start_year: raise ValueError('Năm kết thúc quá trình công tác không được nhỏ hơn năm bắt đầu.')
        obj=l.employment_records.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or EmploymentRecord(lecturer=l)
        obj.start_year=start_year; obj.end_year=end_year; obj.position=position[:250]; obj.organization=organization[:500]; obj.organization_address=organization_address[:500]
        obj.save(); kept.append(obj.pk)
    finish(l.employment_records, kept)

    kept=[]
    for i,x in enumerate(ip_records):
        title=str(x.get('title_content','')).strip()
        if not title: continue
        obj=l.ip_records.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or IntellectualPropertyRecord(lecturer=l)
        obj.title_content=title[:900]; obj.grant_year=year_value(x.get('grant_year'), required=True); obj.sort_order=i; obj.save(); kept.append(obj.pk)
    finish(l.ip_records, kept)

    kept=[]
    for i,x in enumerate(practical):
        work_name=str(x.get('work_name','')).strip(); detail=str(x.get('application_detail','')).strip()
        if not (work_name or detail): continue
        start_date,end_date=date_value(x.get('start_date')),date_value(x.get('end_date'))
        if start_date and end_date and end_date < start_date: raise ValueError('Thời gian áp dụng thực tiễn không hợp lệ.')
        obj=l.practical_applications.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or PracticalApplicationRecord(lecturer=l)
        obj.work_name=work_name[:900]; obj.application_detail=detail[:10000]; obj.start_date=start_date; obj.end_date=end_date; obj.sort_order=i; obj.save(); kept.append(obj.pk)
    finish(l.practical_applications, kept)

    kept=[]
    for i,x in enumerate(tasks):
        name=str(x.get('name','')).strip()
        if not name: continue
        role=x.get('participation_role'); status=x.get('status')
        if role not in ResearchTaskRecord.ParticipationRole.values or status not in ResearchTaskRecord.TaskStatus.values:
            raise ValueError('Vai trò hoặc tình trạng nhiệm vụ không hợp lệ.')
        start_date,end_date=date_value(x.get('start_date')),date_value(x.get('end_date'))
        if start_date and end_date and end_date < start_date: raise ValueError('Thời gian nhiệm vụ KH,CN&ĐMST không hợp lệ.')
        obj=l.research_tasks.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or ResearchTaskRecord(lecturer=l)
        obj.name=name[:900]; obj.participation_role=role; obj.start_date=start_date; obj.end_date=end_date
        obj.program_name=str(x.get('program_name','')).strip()[:500]; obj.status=status; obj.sort_order=i; obj.save(); kept.append(obj.pk)
    finish(l.research_tasks, kept)

    kept=[]
    for i,x in enumerate(awards):
        description=str(x.get('description','')).strip()
        if not description: continue
        obj=l.awards.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or AwardRecord(lecturer=l)
        obj.description=description[:900]; obj.award_year=year_value(x.get('award_year'), required=True); obj.sort_order=i; obj.save(); kept.append(obj.pk)
    finish(l.awards, kept)


@api_view(['GET', 'PATCH'])
def profile(request):
    guard = require_role(request, 'lecturer', 'viewer_admin') if request.method == 'GET' else require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request) if request.method == 'PATCH' else {}
    l = target_lecturer(request, d)
    if not l:
        message = 'Cần chọn giảng viên.' if request.user.role in ('admin','viewer_admin') else 'Tài khoản chưa có hồ sơ giảng viên.'
        return failure(message, 404)
    if request.method == 'PATCH':
        try:
            with transaction.atomic():
                before = profile_payload(l)
                _sync_science_profile(l, d.get('science', {}))
                _sync_profile_lists(l, d)
                after = profile_payload(l)
                audit(request.user, 'update_bm04_profile', 'lecturer', l.id, lecturer=l, before=before, after=after)
        except (ValueError, TypeError) as ex:
            return failure(str(ex))
    return Response(profile_payload(l))


@api_view(['GET', 'POST'])
def works(request):
    if request.method == 'GET':
        qs = ResearchWork.objects.select_related('work_type', 'created_by').prefetch_related('authors', 'evidence')
        if has_role(request.user, 'lecturer'):
            l = own_lecturer(request)
            qs = qs.filter(Q(status='verified') | Q(created_by=request.user) | Q(authors__lecturer=l)).distinct()
        elif not has_role(request.user, 'admin', 'viewer_admin', 'reviewer', 'signer'):
            return failure('Không có quyền.', 403)
        if request.query_params.get('status'):
            qs = qs.filter(status=request.query_params['status'])
        q = request.query_params.get('q', '').strip()
        if q:
            qs = qs.filter(Q(title__icontains=q) | Q(doi__icontains=q))
        return Response([work_out(w) for w in qs[:300]])

    guard = require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request)
    l = target_lecturer(request, d)
    if not l:
        return failure('Admin cần chọn giảng viên.' if request.user.role == 'admin' else 'Chưa có hồ sơ giảng viên.')
    title = str(d.get('title', '')).strip()
    try:
        year = year_value(d.get('publication_year'), required=True)
    except (ValueError, TypeError) as ex:
        return failure(str(ex))
    if not title:
        return failure('Tên công trình là bắt buộc.')
    typ = WorkType.objects.filter(pk=d.get('work_type_id')).first()
    if not typ:
        return failure('Loại công trình không hợp lệ.')
    doi = normalize_doi(d.get('doi', ''))
    dup = find_duplicate(title, year, doi)
    if dup:
        return failure('Công trình có khả năng đã tồn tại; hãy gửi yêu cầu liên kết.', 409, duplicate={'id': str(dup.id), 'title': dup.title, 'status': dup.status})
    authorship = d.get('authorship_type', WorkAuthor.AuthorshipType.AUTHOR)
    if authorship not in WorkAuthor.AuthorshipType.values:
        return failure('Loại tác giả không hợp lệ.')
    with transaction.atomic():
        w = ResearchWork.objects.create(
            title=title, title_norm=normalize_title(title), publication_year=year,
            work_type=typ, doi=doi, doi_norm=doi, venue=str(d.get('venue', ''))[:500], created_by=request.user,
        )
        publication_department = Department.objects.filter(pk=d.get('publication_department_id', l.department_id)).first() or l.department
        WorkAuthor.objects.create(
            work=w, lecturer=l, name=l.full_name, author_order=1,
            authorship_type=authorship, department_at_publication=publication_department,
        )
        authors = d.get('authors', []) if isinstance(d.get('authors', []), list) else []
        for i, a in enumerate(authors[:50], 2):
            name = str(a.get('name', '') if isinstance(a, dict) else a).strip()
            if name:
                WorkAuthor.objects.create(work=w, name=name[:200], author_order=i, authorship_type='coauthor')
        audit(request.user, 'create_work', 'work', w.id)
    return Response(work_out(w), status=201)


@api_view(['GET', 'PATCH'])
def work_detail(request, pk):
    w = ResearchWork.objects.select_related('work_type').prefetch_related('authors', 'evidence').filter(pk=pk).first()
    if not w or not can_access_work(request, w):
        return failure('Không tìm thấy công trình.', 404)
    if request.method == 'PATCH':
        # V8: công trình là dữ liệu thuộc Lý lịch BM-04. Không còn bước gửi/xác minh
        # riêng trước khi tạo BM-04; snapshot lịch sử bảo đảm hồ sơ đã ký không đổi.
        owns_or_authors = w.created_by_id == request.user.id or w.authors.filter(lecturer__user=request.user).exists()
        if request.user.role != 'admin' and not owns_or_authors:
            return failure('Không có quyền sửa công trình này.', 403)
        d = json_body(request)
        try:
            year = year_value(d.get('publication_year', w.publication_year), required=True)
        except (ValueError, TypeError) as ex:
            return failure(str(ex))
        title = str(d.get('title', w.title)).strip()
        doi = normalize_doi(d.get('doi', w.doi))
        if not title:
            return failure('Tên công trình là bắt buộc.')
        dup = find_duplicate(title, year, doi, exclude=w.pk)
        if dup:
            return failure('Trùng công trình đã tồn tại.', 409, duplicate={'id': str(dup.id), 'title': dup.title})
        typ = WorkType.objects.filter(pk=d.get('work_type_id', w.work_type_id)).first()
        if not typ:
            return failure('Loại công trình không hợp lệ.')
        w.title = title
        w.title_norm = normalize_title(title)
        w.doi = doi
        w.doi_norm = doi
        w.venue = str(d.get('venue', w.venue))[:500]
        w.publication_year = year
        w.work_type = typ
        w.status = 'draft'
        w.save()
        audit(request.user, 'update_work', 'work', w.pk)
    return Response(work_out(w))


@api_view(['POST'])
def upload_evidence(request, pk):
    w = ResearchWork.objects.filter(pk=pk).first()
    owns_or_authors = bool(w) and (w.created_by_id == request.user.id or w.authors.filter(lecturer__user=request.user).exists())
    if not w or (request.user.role != 'admin' and not owns_or_authors):
        return failure('Không thể bổ sung minh chứng cho công trình này.', 403)
    f = request.FILES.get('file')
    if not f:
        return failure('Chưa chọn tệp minh chứng công trình.')
    try:
        mode, valid_from, valid_until = _validity_fields(request.data)
        rel, name, sha, size, mime = store_private(f, 'evidence')
    except (DjangoValidationError, ValueError) as ex:
        return error_validation(ex) if isinstance(ex, DjangoValidationError) else failure(str(ex))
    e = WorkEvidence.objects.create(
        work=w, private_path=rel, original_name=name, mime_type=mime, sha256=sha, size=size,
        uploaded_by=request.user, validity_mode=mode, valid_from=valid_from, valid_until=valid_until,
    )
    audit(
        request.user, 'upload_work_evidence', 'work', w.id, file_id=e.id,
        validity_mode=mode, valid_from=str(valid_from or ''), valid_until=str(valid_until or '')
    )
    return Response(evidence_payload(e), status=201)


@api_view(['GET'])
def download_evidence(request, pk):
    e = WorkEvidence.objects.select_related('work').filter(pk=pk).first()
    if not e:
        return failure('Không tìm thấy minh chứng công trình.', 404)
    if not (has_role(request.user, 'admin', 'viewer_admin', 'reviewer', 'signer') or e.work.created_by_id == request.user.id or e.work.authors.filter(lecturer__user=request.user).exists()):
        return failure('Không có quyền tải minh chứng công trình.', 403)
    audit(request.user, 'download_evidence', 'evidence', e.id)
    path = read_private(e.private_path)
    if not path.is_file():
        raise Http404
    response = FileResponse(path.open('rb'), as_attachment=True, filename=e.original_name, content_type='application/octet-stream')
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@api_view(['GET'])
def view_evidence(request, pk):
    e = WorkEvidence.objects.select_related('work').filter(pk=pk).first()
    if not e:
        return failure('Không tìm thấy minh chứng công trình.', 404)
    if not (has_role(request.user, 'admin', 'viewer_admin', 'reviewer', 'signer') or e.work.created_by_id == request.user.id or e.work.authors.filter(lecturer__user=request.user).exists()):
        return failure('Không có quyền xem minh chứng công trình.', 403)
    path = read_private(e.private_path)
    if not path.is_file():
        raise Http404
    mime = e.mime_type if e.mime_type in ('application/pdf','image/png','image/jpeg') else 'application/octet-stream'
    audit(request.user, 'view_evidence', 'evidence', e.id)
    response = FileResponse(path.open('rb'), as_attachment=False, filename=e.original_name, content_type=mime)
    response['Content-Disposition'] = f'inline; filename="{e.original_name.replace(chr(34), "")}"'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


def _resolve_evidence_entity(lecturer, section, entity_id):
    entity_id = str(entity_id or '')
    if section == EvidenceDocument.Section.EDUCATION:
        return lecturer.education_records.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.EMPLOYMENT:
        return lecturer.employment_records.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.WORK:
        return ResearchWork.objects.filter(pk=entity_id, authors__lecturer=lecturer).distinct().first()
    if section == EvidenceDocument.Section.IP:
        return lecturer.ip_records.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.PRACTICAL:
        return lecturer.practical_applications.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.TASK:
        return lecturer.research_tasks.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.AWARD:
        return lecturer.awards.filter(pk=entity_id).first()
    if section == EvidenceDocument.Section.ACHIEVEMENT:
        return ScienceProfile.objects.filter(lecturer=lecturer).first() if entity_id in ('', 'science') else None
    return None


def _can_access_profile_evidence(user, lecturer):
    return has_role(user, 'admin', 'viewer_admin', 'reviewer', 'signer') or lecturer.user_id == user.id


def _evidence_queryset(lecturer, section=None, entity_id=None):
    qs = EvidenceDocument.objects.filter(lecturer=lecturer, is_active=True)
    if section:
        if entity_id is None:
            qs = qs.filter(Q(links__section=section) | Q(section=section))
        else:
            target = str(entity_id)
            qs = qs.filter(Q(links__section=section, links__entity_id=target) | Q(section=section, entity_id=target))
    return qs.distinct().order_by('created_at', 'id')


def _evidence_out(e):
    data = evidence_payload(e)
    data.update({
        'section': e.section, 'section_label': e.get_section_display(), 'entity_id': e.entity_id,
        'uploaded_by': e.uploaded_by.username,
        'links': [{'section': x.section, 'section_label': x.get_section_display(), 'entity_id': x.entity_id} for x in e.links.all()],
    })
    return data


def _validity_fields(data):
    """Đọc lựa chọn thời hạn minh chứng.

    V8 tạm khóa chức năng hiệu lực theo ngày. Hệ thống vẫn lưu lựa chọn
    `indefinite`/`fixed`, nhưng khi chọn `fixed` sẽ không yêu cầu và không
    sử dụng valid_from/valid_until. Hai trường ngày được để trống để tránh
    vô tình kích hoạt logic hết hạn cũ.
    """
    mode = str(data.get('validity_mode') or EvidenceDocument.ValidityMode.INDEFINITE)
    if mode not in EvidenceDocument.ValidityMode.values:
        raise ValueError('Thời hạn hiệu lực không hợp lệ.')
    return mode, None, None


@api_view(['GET', 'POST'])
def profile_evidence(request):
    """Minh chứng xác minh cho dữ liệu khoa học.

    POST có hai chế độ:
    - upload file mới: gửi file + section + entity_id + validity_mode...
    - tái sử dụng file đã có: gửi existing_evidence_id + section + entity_id.
    """
    if request.method == 'GET':
        lecturer_id = request.query_params.get('lecturer_id')
        lecturer = Lecturer.objects.filter(pk=lecturer_id).first() if lecturer_id else own_lecturer(request)
        if not lecturer or not _can_access_profile_evidence(request.user, lecturer):
            return failure('Không có quyền xem minh chứng.', 403)
        section = request.query_params.get('section')
        entity_id = request.query_params.get('entity_id')
        qs = _evidence_queryset(lecturer, section, entity_id)
        return Response([_evidence_out(e) for e in qs.select_related('uploaded_by').prefetch_related('links')[:500]])

    lecturer_id = request.data.get('lecturer_id')
    lecturer = Lecturer.objects.filter(pk=lecturer_id).first() if lecturer_id else own_lecturer(request)
    if not lecturer:
        return failure('Không tìm thấy giảng viên.', 404)
    if request.user.role == 'lecturer' and lecturer.user_id != request.user.id:
        return failure('Không được tải minh chứng cho giảng viên khác.', 403)
    if request.user.role not in ('lecturer', 'admin') and not request.user.is_superuser:
        return failure('Không có quyền tải minh chứng.', 403)
    section = str(request.data.get('section') or '')
    entity_id = str(request.data.get('entity_id') or '')
    if section not in EvidenceDocument.Section.values:
        return failure('Nhóm minh chứng không hợp lệ.')
    if section == EvidenceDocument.Section.WORK:
        return failure('Minh chứng công trình khoa học được quản lý trực tiếp tại mục Công trình khoa học.', 409)
    if not _resolve_evidence_entity(lecturer, section, entity_id):
        return failure('Không tìm thấy dữ liệu cần gắn minh chứng.', 404)

    existing_id = request.data.get('existing_evidence_id')
    if existing_id:
        e = EvidenceDocument.objects.filter(pk=existing_id, lecturer=lecturer, is_active=True).first()
        if not e:
            return failure('Không tìm thấy minh chứng đã có.', 404)
        EvidenceLink.objects.get_or_create(evidence=e, section=section, entity_id=entity_id, defaults={'created_by': request.user})
        audit(request.user, 'link_profile_evidence', 'evidence_document', e.pk, lecturer=lecturer,
              target_section=section, target_entity_id=entity_id, file=e.original_name)
        return Response(_evidence_out(EvidenceDocument.objects.prefetch_related('links').select_related('uploaded_by').get(pk=e.pk)))

    f = request.FILES.get('file')
    if not f:
        return failure('Chưa chọn tệp minh chứng.')
    try:
        mode, valid_from, valid_until = _validity_fields(request.data)
        rel, name, sha, size, mime = store_private(f, 'profile_evidence')
    except (DjangoValidationError, ValueError) as ex:
        if isinstance(ex, DjangoValidationError):
            return error_validation(ex)
        return failure(str(ex))
    e = EvidenceDocument.objects.create(
        lecturer=lecturer, section=section, entity_id=entity_id, validity_mode=mode,
        valid_from=valid_from, valid_until=valid_until, private_path=rel,
        original_name=name, mime_type=mime, sha256=sha, size=size, uploaded_by=request.user,
    )
    EvidenceLink.objects.create(evidence=e, section=section, entity_id=entity_id, created_by=request.user)
    audit(request.user, 'upload_profile_evidence', 'evidence_document', e.pk, lecturer=lecturer,
          target_section=section, target_entity_id=entity_id, file=name, sha256=sha,
          validity_mode=mode, valid_from=str(valid_from or ''), valid_until=str(valid_until or ''))
    return Response(_evidence_out(EvidenceDocument.objects.prefetch_related('links').select_related('uploaded_by').get(pk=e.pk)), status=201)


@api_view(['GET', 'DELETE'])
def profile_evidence_detail(request, pk):
    e = EvidenceDocument.objects.select_related('lecturer', 'uploaded_by').prefetch_related('links').filter(pk=pk).first()
    if not e or not _can_access_profile_evidence(request.user, e.lecturer):
        return failure('Không tìm thấy minh chứng.', 404)
    if request.method == 'DELETE':
        if request.user.role == 'lecturer' and e.lecturer.user_id != request.user.id:
            return failure('Không có quyền.', 403)
        if request.user.role not in ('lecturer', 'admin') and not request.user.is_superuser:
            return failure('Không có quyền.', 403)
        e.is_active = False
        e.save(update_fields=['is_active'])
        audit(request.user, 'deactivate_profile_evidence', 'evidence_document', e.pk, lecturer=e.lecturer,
              target_section=e.section, target_entity_id=e.entity_id, file=e.original_name)
        return Response({'ok': True})
    return Response(_evidence_out(e))


def _profile_evidence_file(request, pk, inline):
    e = EvidenceDocument.objects.select_related('lecturer').filter(pk=pk).first()
    if not e or not _can_access_profile_evidence(request.user, e.lecturer):
        return failure('Không tìm thấy minh chứng.', 404)
    path = read_private(e.private_path)
    if not path.is_file():
        raise Http404
    mime = e.mime_type if e.mime_type in ('application/pdf','image/png','image/jpeg') else 'application/octet-stream'
    audit(request.user, 'view_profile_evidence' if inline else 'download_profile_evidence', 'evidence_document', e.pk,
          lecturer=e.lecturer, target_section=e.section, target_entity_id=e.entity_id)
    resp = FileResponse(path.open('rb'), as_attachment=not inline, filename=e.original_name, content_type=mime)
    if inline:
        resp['Content-Disposition'] = f'inline; filename="{e.original_name.replace(chr(34), "")}"'
    resp['X-Content-Type-Options'] = 'nosniff'
    return resp


@api_view(['GET'])
def view_profile_evidence(request, pk):
    return _profile_evidence_file(request, pk, True)


@api_view(['GET'])
def download_profile_evidence(request, pk):
    return _profile_evidence_file(request, pk, False)


@api_view(['POST'])
def submit_work(request, pk):
    return failure(
        'V8 không còn quy trình gửi xác minh công trình riêng. Hãy cập nhật công trình tại Lý lịch BM-04; Phòng NCKH chỉ xác minh khi công trình được đưa vào hồ sơ BM-04 và gửi phê duyệt.',
        409,
    )


@api_view(['POST'])
def review_work(request, pk):
    return failure(
        'V8 không còn xác minh công trình độc lập. Việc xác minh được thực hiện trong checklist của hồ sơ BM-04 tại In & phê duyệt BM-04.',
        409,
    )


@api_view(['GET', 'POST'])
def link_requests(request):
    if request.method == 'GET':
        guard = require_role(request, 'reviewer', 'admin')
        if guard:
            return guard
        qs = WorkLinkRequest.objects.select_related('work', 'lecturer', 'department_at_publication')
        return Response([{
            'id': x.id, 'work_id': str(x.work_id), 'work_title': x.work.title,
            'lecturer': x.lecturer.full_name, 'authorship_type': x.requested_authorship_type,
            'department_at_publication': x.department_at_publication.name if x.department_at_publication else '',
            'status': x.status,
        } for x in qs.filter(status='pending')[:100]])
    guard = require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request)
    w = ResearchWork.objects.filter(pk=d.get('work_id'), status='verified').first()
    l = own_lecturer(request)
    if not w or not l:
        return failure('Không tìm thấy công trình đã xác minh.')
    if WorkAuthor.objects.filter(work=w, lecturer=l).exists():
        return failure('Công trình đã có trong hồ sơ của bạn.')
    authorship = d.get('authorship_type', 'coauthor')
    if authorship not in WorkAuthor.AuthorshipType.values:
        return failure('Loại tác giả không hợp lệ.')
    dep = Department.objects.filter(pk=d.get('publication_department_id', l.department_id)).first()
    if not dep:
        return failure('Khoa tại thời điểm công bố không hợp lệ.')
    x, created = WorkLinkRequest.objects.get_or_create(
        work=w, lecturer=l,
        defaults={'requested_authorship_type': authorship, 'department_at_publication': dep},
    )
    if not created and x.status != 'pending':
        return failure('Yêu cầu liên kết đã được xử lý.')
    if created:
        audit(request.user, 'create_link', 'work', w.id)
    return Response({'id': x.id, 'status': x.status}, status=201 if created else 200)


@api_view(['POST'])
def review_link(request, pk):
    guard = require_role(request, 'reviewer')
    if guard:
        return guard
    d = json_body(request)
    action = d.get('action')
    if action not in ('approved', 'rejected'):
        return failure('Lựa chọn không hợp lệ.')
    with transaction.atomic():
        x = WorkLinkRequest.objects.select_for_update().select_related('lecturer', 'work').filter(pk=pk).first()
        if not x or x.status != 'pending':
            return failure('Yêu cầu không tồn tại hoặc đã xử lý.', 409)
        if action == 'approved':
            if not x.work.authors.filter(name__iexact=x.lecturer.full_name).exists():
                return failure('Không tìm thấy tác giả trong danh sách đã xác minh; cần kiểm tra công trình trước khi liên kết.', 409)
            a = x.work.authors.filter(name__iexact=x.lecturer.full_name, lecturer__isnull=True).first()
            if not a:
                return failure('Tên tác giả đã được liên kết hoặc cần xử lý thủ công.', 409)
            a.lecturer = x.lecturer
            a.authorship_type = x.requested_authorship_type
            a.department_at_publication = x.department_at_publication
            a.save(update_fields=['lecturer', 'authorship_type', 'department_at_publication'])
        x.status = action
        x.reviewer = request.user
        x.reviewed_at = timezone.now()
        x.save()
        audit(request.user, 'review_link', 'work', x.work_id, result=action)
    return Response({'ok': True, 'status': action})


@api_view(['GET', 'POST'])
def corrections(request):
    if request.method == 'GET':
        guard = require_role(request, 'reviewer', 'admin')
        if guard:
            return guard
        return Response([{'id': c.id, 'work_id': str(c.work_id), 'work_title': c.work.title, 'proposed': c.proposed, 'status': c.status} for c in WorkCorrection.objects.select_related('work').filter(status='pending')[:100]])
    guard = require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request)
    w = ResearchWork.objects.filter(pk=d.get('work_id'), status='verified').first()
    if not w or not w.authors.filter(lecturer=own_lecturer(request)).exists():
        return failure('Không có quyền sửa công trình.', 403)
    proposed = d.get('proposed', {})
    allowed = {'title', 'publication_year', 'doi', 'venue', 'work_type_id'}
    if not isinstance(proposed, dict) or not proposed or not set(proposed).issubset(allowed):
        return failure('Dữ liệu chỉnh sửa không hợp lệ.')
    c = WorkCorrection.objects.create(work=w, proposer=request.user, proposed=proposed)
    audit(request.user, 'create_correction', 'work', w.id)
    return Response({'id': c.id, 'status': c.status}, status=201)


@api_view(['POST'])
def review_correction(request, pk):
    guard = require_role(request, 'reviewer')
    if guard:
        return guard
    d = json_body(request)
    action = d.get('action')
    note = str(d.get('comment', ''))
    if action not in ('approved', 'rejected'):
        return failure('Hành động không hợp lệ.')
    with transaction.atomic():
        c = WorkCorrection.objects.select_for_update().select_related('work').filter(pk=pk).first()
        if not c or c.status != 'pending':
            return failure('Yêu cầu không tồn tại hoặc đã xử lý.', 409)
        w = ResearchWork.objects.select_for_update().get(pk=c.work_id)
        if action == 'approved':
            p = c.proposed
            title = str(p.get('title', w.title)).strip()
            try:
                year = year_value(p.get('publication_year', w.publication_year), required=True)
            except (ValueError, TypeError) as ex:
                return failure(str(ex))
            doi = normalize_doi(p.get('doi', w.doi))
            if find_duplicate(title, year, doi, exclude=w.pk):
                return failure('Chỉnh sửa gây trùng công trình.', 409)
            typ = WorkType.objects.filter(pk=p.get('work_type_id', w.work_type_id)).first()
            if not typ:
                return failure('Loại công trình không hợp lệ.')
            WorkVersion.objects.create(work=w, previous={'title': w.title, 'year': w.publication_year, 'doi': w.doi, 'venue': w.venue, 'work_type_id': w.work_type_id})
            w.title = title
            w.title_norm = normalize_title(title)
            w.publication_year = year
            w.doi = doi
            w.doi_norm = doi
            w.venue = str(p.get('venue', w.venue))[:500]
            w.work_type = typ
            w.save()
        c.status = action
        c.comment = note
        c.reviewer = request.user
        c.reviewed_at = timezone.now()
        c.save()
        audit(request.user, 'review_correction', 'work', w.pk, result=action)
    return Response({'ok': True, 'status': action})


def _selected_ids(d, key, *, max_items=200):
    raw = d.get(key, [])
    if raw in (None, ''):
        return []
    if not isinstance(raw, list):
        raise ValueError(f'{key} phải là một danh sách.')
    result = []
    for value in raw:
        value = str(value).strip()
        if value and value not in result:
            result.append(value)
    if len(result) > max_items:
        raise ValueError(f'{key} vượt quá số mục cho phép.')
    return result



def _bm04_missing_required(l):
    checks = [
        ('Số định danh cá nhân', l.personal_id), ('Ngày sinh', l.date_of_birth), ('Giới tính', l.gender),
        ('Địa chỉ nhà riêng', l.home_address), ('Điện thoại liên hệ', l.phone), ('E-mail', l.email),
        ('Địa chỉ đơn vị', l.department.address), ('Thủ trưởng đơn vị', l.department.head_name),
        ('Điện thoại đơn vị', l.department.phone),
    ]
    return [label for label, value in checks if not value]

def _application_payload(d):
    role = d.get('registered_role')
    if role not in ProfileApproval.RegisteredRole.values:
        raise ValueError('Phải chọn chức danh đăng ký thực hiện nhiệm vụ.')
    task_name = str(d.get('task_name', '')).strip()
    if not task_name:
        raise ValueError('Tên nhiệm vụ là bắt buộc theo BM-04.')
    other = str(d.get('registered_role_other', '')).strip()
    if role == 'other' and not other:
        raise ValueError('Cần ghi rõ chức danh khác trong Ban Chủ nhiệm nhiệm vụ.')
    work_ids = _selected_ids(d, 'selected_work_ids', max_items=5)
    return {
        'task_name': task_name[:900], 'registered_role': role,
        'registered_role_other': other[:250], 'signing_place': str(d.get('signing_place', '')).strip()[:180],
        'signing_date': date_value(d.get('signing_date')) or timezone.localdate(),
        'selected_work_ids': work_ids,
        'selected_ip_ids': _selected_ids(d, 'selected_ip_ids'),
        'selected_application_ids': _selected_ids(d, 'selected_application_ids'),
        'selected_award_ids': _selected_ids(d, 'selected_award_ids'),
    }


def _validate_application_ownership(l, app):
    work_count = WorkAuthor.objects.filter(lecturer=l, work_id__in=app['selected_work_ids']).values('work_id').distinct().count()
    if work_count != len(set(app['selected_work_ids'])):
        raise ValueError('Danh sách công trình có mục không thuộc hồ sơ giảng viên.')
    checks = [
        (l.ip_records, app['selected_ip_ids'], 'văn bằng sở hữu trí tuệ'),
        (l.practical_applications, app['selected_application_ids'], 'công trình áp dụng thực tiễn'),
        (l.awards, app['selected_award_ids'], 'giải thưởng'),
    ]
    for manager, ids, label in checks:
        if manager.filter(pk__in=ids).count() != len(set(ids)):
            raise ValueError(f'Danh sách {label} có dữ liệu không thuộc hồ sơ.')


@api_view(['GET', 'POST'])
def approvals(request):
    if request.method == 'GET':
        qs = ProfileApproval.objects.select_related('lecturer', 'lecturer__department').order_by('-submitted_at')
        if has_role(request.user, 'lecturer'):
            l = own_lecturer(request)
            qs = qs.filter(lecturer=l) if l else qs.none()
        elif not has_role(request.user, 'reviewer', 'signer', 'admin', 'viewer_admin'):
            return failure('Không có quyền.', 403)

        q = str(request.query_params.get('q') or '').strip()
        status_filter = str(request.query_params.get('status') or '').strip()
        lecturer_filter = str(request.query_params.get('lecturer') or '').strip()
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(task_name__icontains=q))
        if status_filter:
            if status_filter not in ProfileApproval.Status.values:
                return failure('Trạng thái BM-04 không hợp lệ.')
            qs = qs.filter(status=status_filter)
        if lecturer_filter:
            qs = qs.filter(
                Q(lecturer__full_name__icontains=lecturer_filter)
                | Q(lecturer__staff_code__icontains=lecturer_filter)
            )

        # Giữ tương thích với Dashboard cũ: chỉ trả payload phân trang khi client truyền page.
        if request.query_params.get('page') is not None:
            try:
                page_number = max(1, int(request.query_params.get('page', '1')))
                page_size = int(request.query_params.get('page_size', '10'))
            except ValueError:
                return failure('Tham số phân trang không hợp lệ.')
            page_size = min(max(page_size, 5), 100)
            paginator = Paginator(qs, page_size)
            page_obj = paginator.get_page(page_number)
            return Response({
                'results': [approval_out(a) for a in page_obj.object_list],
                'count': paginator.count,
                'page': page_obj.number,
                'pages': paginator.num_pages or 1,
                'page_size': page_size,
            })

        return Response([approval_out(a) for a in qs[:200]])
    guard = require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request)
    l = target_lecturer(request, d)
    if not l:
        return failure('Admin cần chọn giảng viên.' if request.user.role == 'admin' else 'Chưa có hồ sơ giảng viên.')
    missing = _bm04_missing_required(l)
    if missing:
        return failure('Hồ sơ BM-04 còn thiếu thông tin bắt buộc: ' + ', '.join(missing) + '. Liên hệ Admin để cập nhật dữ liệu nhân sự/đơn vị.')
    try:
        app = _application_payload(d)
        _validate_application_ownership(l, app)
    except (ValueError, TypeError) as ex:
        return failure(str(ex))
    s = snapshot_for(l, app)
    parent = None
    version = 1
    parent_id = d.get('parent_approval_id')
    if parent_id:
        parent = ProfileApproval.objects.filter(pk=parent_id, lecturer=l).first()
        if not parent or parent.status not in (ProfileApproval.Status.CHANGES, ProfileApproval.Status.REJECTED):
            return failure('Chỉ có thể tạo phiên bản mới từ hồ sơ đã bị trả sửa hoặc từ chối.', 409)
        if parent.child_versions.exists():
            return failure('Hồ sơ này đã có phiên bản thay thế.', 409)
        version = parent.version + 1
        base_code = parent.code.split('-V')[0]
        code = f'{base_code}-V{version}'
    else:
        code = f'BM04-{timezone.localdate():%Y%m%d}-{uuid.uuid4().hex[:8].upper()}'

    # Giảng viên phải ký số ngay tại thời điểm gửi yêu cầu xác minh lý lịch.
    # Admin có thể tạo hồ sơ thay giảng viên theo quyền quản trị, nhưng không được giả mạo
    # chữ ký cá nhân của giảng viên.
    personal_identity = None
    personal_meta = None
    personal_signature_b64 = None
    if request.user.role == 'lecturer':
        if l.user_id != request.user.id:
            return failure('Không được ký BM-04 của giảng viên khác.', 403)
        password = str(d.get('password') or '')
        if not password or not request.user.check_password(password):
            return failure('Mật khẩu xác nhận ký số không đúng.', 403)
        personal_identity = SigningIdentity.objects.filter(user=request.user, is_active=True).first()
        if not personal_identity or not personal_identity.encrypted_private_key or not personal_identity.certificate_pem:
            return failure('Tài khoản chưa được cấp danh tính ký số nội bộ. Hãy tải ảnh chữ ký trước khi gửi BM-04.')
        if not personal_identity.signature_image_path:
            return failure('Chưa tải ảnh chữ ký cá nhân để hiển thị trên BM-04.')
        signed_at = timezone.localtime()
        verification_code = f'SIG-P-{signed_at:%Y%m%d}-{secrets.token_hex(6).upper()}'
        personal_meta = {
            'image_path': str(read_private(personal_identity.signature_image_path)),
            'image_rel_path': personal_identity.signature_image_path,
            'signer_name': personal_identity.display_name or l.full_name,
            'signer_position': personal_identity.position or l.professional_title or l.position,
            'verification_code': verification_code,
            'signed_at': signed_at.strftime('%d/%m/%Y %H:%M:%S'),
        }

    pdf = profile_pdf(s, code, draft=False, personal_signature=personal_meta)
    rel, sha = store_bytes(pdf, 'approvals')
    if personal_identity:
        personal_signature_b64 = sign_bytes(personal_identity, pdf)

    with transaction.atomic():
        a = ProfileApproval.objects.create(
            parent_approval=parent, version=version,
            lecturer=l, code=code, task_name=app['task_name'], registered_role=app['registered_role'],
            registered_role_other=app['registered_role_other'], signing_place=app['signing_place'],
            signing_date=app['signing_date'], snapshot=s, snapshot_sha256=snapshot_hash(s),
            pdf_path=rel, pdf_sha256=sha,
        )
        create_approval_verification_items(a)
        if personal_identity and personal_meta and personal_signature_b64:
            DigitalSignature.objects.create(
                approval=a, signer=request.user, identity=personal_identity,
                signature_role=DigitalSignature.SignatureRole.PERSONAL,
                visible_signature=personal_meta,
                signature_b64=personal_signature_b64, certificate_pem=personal_identity.certificate_pem,
                source_pdf_sha256=sha, signed_pdf_sha256=sha, signed_file_path=rel,
                verification_code=personal_meta['verification_code'], verification_status='valid',
            )
            ApprovalEvent.objects.create(approval=a, actor=request.user, action='personal_digital_signed', note=personal_meta['verification_code'])
            audit(request.user, 'personal_digital_sign_approval', 'approval', a.pk, verification_code=personal_meta['verification_code'], signed_pdf_sha256=sha)
        ApprovalEvent.objects.create(approval=a, actor=request.user, action='submitted')
        audit(request.user, 'submit_approval', 'approval', a.pk, lecturer=l, version=version, parent_approval_id=str(parent.pk) if parent else None)
    return Response(approval_out(a), status=201)


@api_view(['POST'])
def draft_pdf(request):
    guard = require_role(request, 'lecturer')
    if guard:
        return guard
    d = json_body(request)
    l = target_lecturer(request, d)
    if not l:
        return failure('Admin cần chọn giảng viên.' if request.user.role == 'admin' else 'Chưa có hồ sơ.')
    try:
        app = _application_payload(d)
        _validate_application_ownership(l, app)
    except (ValueError, TypeError) as ex:
        return failure(str(ex))
    data = profile_pdf(snapshot_for(l, app), 'BAN-NHAP-BM04', draft=True)
    from io import BytesIO
    resp = FileResponse(BytesIO(data), as_attachment=True, filename='BM-04-ly-lich-khoa-hoc-ban-nhap.pdf', content_type='application/pdf')
    resp['X-Content-Type-Options'] = 'nosniff'
    return resp


@api_view(['GET'])
def approval_detail(request, pk):
    a = ProfileApproval.objects.select_related('lecturer', 'lecturer__department').filter(pk=pk).first()
    if not a:
        return failure('Không thấy hồ sơ.', 404)
    if has_role(request.user, 'lecturer') and a.lecturer.user_id != request.user.id:
        return failure('Không có quyền.', 403)

    # Chi tiết BM-04 luôn đọc từ snapshot bất biến, không lấy dữ liệu công trình hiện tại.
    # Nhờ đó mở lại BM-04 cũ sẽ thấy đúng thông tin/minh chứng tại thời điểm đã gửi.
    snapshot_works = a.snapshot.get('works', []) if isinstance(a.snapshot, dict) else []
    linked_works = [{
        'id': str(x.get('id') or ''), 'title': x.get('title') or '', 'work_type': x.get('type') or '',
        'publication_year': x.get('year'), 'doi': x.get('doi') or '', 'venue': x.get('venue') or '',
        'status': 'snapshot', 'authors': x.get('authors') or [],
        # V6.4: WorkEvidence chính là minh chứng công trình khoa học bắt buộc.
        # Snapshot V6.3 đã lưu chúng dưới key `attachments`, nên dùng fallback
        # để hồ sơ cũ đang xử lý vẫn xem/xác minh được đúng file đã chốt.
        'evidence': x.get('verification_evidence') or x.get('attachments') or [],
        'verification_evidence': x.get('verification_evidence') or x.get('attachments') or [],
        'authorship_label': x.get('authorship_label') or '',
    } for x in snapshot_works if isinstance(x, dict)]

    def approval_item_evidence(item):
        evidence = item.evidence_snapshot or []
        if evidence or item.section != EvidenceDocument.Section.WORK:
            return evidence
        target = str(item.entity_id or '')
        for work in snapshot_works:
            if isinstance(work, dict) and str(work.get('id') or '') == target:
                return work.get('verification_evidence') or work.get('attachments') or []
        return []

    return Response({
        **approval_out(a),
        'lecturer_code': a.lecturer.staff_code,
        'lecturer_department': a.lecturer.department.name,
        'snapshot': a.snapshot,
        'linked_works': linked_works,
        'verification_items': [{
            'id': x.pk, 'section': x.section, 'section_label': x.get_section_display(),
            'entity_id': x.entity_id, 'label': x.label, 'status': x.status,
            'evidence_required': x.evidence_required, 'evidence': approval_item_evidence(x),
            'reused': bool(x.reused_verification_id), 'comment': x.comment,
            'reviewer': x.reviewer.username if x.reviewer else None,
            'reviewed_at': x.reviewed_at.isoformat() if x.reviewed_at else None,
        } for x in a.verification_items.select_related('reviewer', 'reused_verification')],
        'lecturer_logs': [{
            'at': x.at.isoformat(), 'actor': x.actor.username if x.actor else None,
            'action': x.action, 'entity': x.entity, 'entity_id': x.entity_id, 'details': x.details,
        } for x in AuditLog.objects.select_related('actor').filter(lecturer=a.lecturer)[:100]],
        'events': [{'action': e.action, 'note': e.note, 'actor': e.actor.username, 'at': e.at.isoformat()} for e in a.events.select_related('actor')],
    })


@api_view(['POST'])
def review_approval_item(request, pk, item_pk):
    """Phòng NCKH xác minh từng mục trực tiếp từ checklist của BM-04."""
    guard = require_role(request, 'reviewer')
    if guard:
        return guard
    d = json_body(request)
    action = str(d.get('action') or '')
    comment = str(d.get('comment') or '').strip()
    with transaction.atomic():
        a = ProfileApproval.objects.select_for_update().filter(pk=pk).first()
        if not a:
            return failure('Không thấy hồ sơ.', 404)
        if a.status != ProfileApproval.Status.REVIEW:
            return failure('Chỉ xác minh từng mục khi hồ sơ đang chờ NCKH.', 409)
        item = ApprovalVerificationItem.objects.select_for_update().filter(pk=item_pk, approval=a).first()
        if not item:
            return failure('Không thấy mục cần xác minh.', 404)
        try:
            review_item(item, request.user, action, comment)
        except ValueError as ex:
            return failure(str(ex))
        ApprovalEvent.objects.create(
            approval=a, actor=request.user, action=f'verify_item_{action}',
            note=f'{item.get_section_display()}: {item.label}' + (f' — {comment}' if comment else ''),
        )
        audit(request.user, 'review_approval_item', 'approval_verification_item', item.pk,
              lecturer=a.lecturer, approval_id=str(a.pk), section=item.section, target_entity_id=item.entity_id,
              result=action, comment=comment)
    return Response({'ok': True, 'status': item.status, 'verification_complete': all_items_verified(a)})


@api_view(['POST'])
def review_approval(request, pk):
    guard = require_role(request, 'reviewer', 'signer')
    if guard:
        return guard
    d = json_body(request)
    action = d.get('action')
    note = str(d.get('comment', '')).strip()
    with transaction.atomic():
        a = ProfileApproval.objects.select_for_update().filter(pk=pk).first()
        if not a:
            return failure('Không thấy hồ sơ.', 404)
        if request.user.role == 'admin' or request.user.is_superuser:
            transitions = {
                'review': {'forward': 'signer', 'changes': 'changes', 'reject': 'rejected'},
                'signer': {'approve': 'print', 'changes': 'changes', 'reject': 'rejected'},
                'print': {'signed': 'signed'},
            }
        elif has_role(request.user, 'reviewer'):
            transitions = {'review': {'forward': 'signer', 'changes': 'changes', 'reject': 'rejected'}, 'print': {'signed': 'signed'}}
        else:
            transitions = {'signer': {'approve': 'print', 'changes': 'changes', 'reject': 'rejected'}}
        next_state = transitions.get(a.status, {}).get(action)
        if not next_state:
            return failure('Hành động không hợp lệ với vai trò hoặc trạng thái hiện tại.', 409)
        if a.status == ProfileApproval.Status.REVIEW and action == 'forward' and not all_items_verified(a):
            pending = a.verification_items.exclude(status=ApprovalVerificationItem.Status.VERIFIED).count()
            return failure(f'Còn {pending} mục chưa xác minh. Cần hoàn tất checklist trước khi chuyển lãnh đạo.', 409)
        if action in ('changes', 'reject') and not note:
            return failure('Phải nhập lý do.')
        if has_role(request.user, 'reviewer') or (request.user.role == 'admin' and a.status in ('review', 'print')):
            a.reviewer = request.user
            a.reviewed_at = timezone.now()
        if has_role(request.user, 'signer') or (request.user.role == 'admin' and a.status == 'signer'):
            a.signer = request.user
            a.approved_at = timezone.now()
        a.status = next_state
        a.comment = note
        a.save()
        ApprovalEvent.objects.create(approval=a, actor=request.user, action=action, note=note)
        audit(request.user, 'review_approval', 'approval', a.pk, lecturer=a.lecturer, result=next_state, comment=note)
    return Response(approval_out(a))


@api_view(['POST'])
def upload_scan(request, pk):
    guard = require_role(request, 'reviewer')
    if guard:
        return guard
    file = request.FILES.get('file')
    if not file:
        return failure('Chưa chọn bản scan.')
    with transaction.atomic():
        a = ProfileApproval.objects.select_for_update().filter(pk=pk).first()
        if not a:
            return failure('Không có hồ sơ.', 404)
        if a.status != 'signed':
            return failure('Phải xác nhận đã ký tay trước khi tải bản scan.', 409)
        try:
            rel, name, sha, size, mime = store_private(file, 'signed_scans', extensions=('.pdf',))
        except DjangoValidationError as ex:
            return error_validation(ex)
        a.scan_path = rel
        a.scan_sha256 = sha
        a.status = 'completed'
        a.completed_at = timezone.now()
        a.save()
        ApprovalEvent.objects.create(approval=a, actor=request.user, action='completed', note=name)
        audit(request.user, 'complete_approval', 'approval', a.pk, scan_hash=sha)
    return Response(approval_out(a))


@api_view(['GET'])
def approval_file(request, pk, kind):
    a = ProfileApproval.objects.select_related('lecturer').filter(pk=pk).first()
    if not a:
        return failure('Không tìm thấy hồ sơ.', 404)
    if has_role(request.user, 'lecturer') and a.lecturer.user_id != request.user.id:
        return failure('Không có quyền.', 403)
    if kind == 'scan':
        if a.status != 'completed' or not a.scan_path:
            return failure('Bản scan chưa sẵn sàng.', 404)
        rel = a.scan_path
    elif kind == 'signed':
        if not a.signed_pdf_path:
            return failure('Bản ký số chưa sẵn sàng.', 404)
        rel = a.signed_pdf_path
    elif kind == 'pdf':
        rel = a.pdf_path
    else:
        return failure('Loại tệp không hợp lệ.', 404)
    path = read_private(rel)
    if not path.is_file():
        raise Http404
    inline = request.query_params.get('inline') == '1'
    audit(request.user, 'view_approval' if inline else 'download_approval', 'approval', a.pk, kind=kind)
    resp = FileResponse(path.open('rb'), as_attachment=not inline, filename=f'{a.code}-{kind}.pdf', content_type='application/pdf')
    if inline:
        resp['Content-Disposition'] = f'inline; filename="{a.code}-{kind}.pdf"'
    resp['X-Content-Type-Options'] = 'nosniff'
    return resp



@api_view(['GET', 'POST'])
def signing_identity(request):
    """Provision/view an internal signing identity and upload the visible signature image."""
    if request.method == 'GET':
        guard = require_role(request, 'lecturer', 'signer', 'admin', 'viewer_admin')
        if guard:
            return guard
        target = request.user
        user_id = request.query_params.get('user_id')
        if user_id and request.user.role in ('admin', 'viewer_admin'):
            target = User.objects.filter(pk=user_id).first()
            if not target:
                return failure('Không thấy tài khoản người ký.', 404)
        identity = SigningIdentity.objects.filter(user=target).first()
        return Response({'user_id': target.id, 'username': target.username, **identity_public_payload(identity)})

    guard = require_role(request, 'lecturer', 'signer', 'admin')
    if guard:
        return guard
    target = request.user
    user_id = request.data.get('user_id')
    if user_id and request.user.role == 'admin':
        target = User.objects.filter(pk=user_id).first()
        if not target:
            return failure('Không thấy tài khoản người ký.', 404)
    if target.role not in ('lecturer', 'signer', 'admin') and not target.is_superuser:
        return failure('Chỉ giảng viên, lãnh đạo ký hoặc quản trị mới được cấp danh tính ký số.', 400)
    image = request.FILES.get('file')
    if not image:
        return failure('Chưa chọn ảnh chữ ký PNG/JPG.')
    try:
        rel, name, sha, size, mime = store_private(image, 'signature_images', extensions=('.png', '.jpg', '.jpeg'))
        identity, _ = SigningIdentity.objects.get_or_create(user=target)
        lecturer = Lecturer.objects.filter(user=target).first()
        default_name = lecturer.full_name if lecturer else (target.get_full_name() or target.username)
        default_position = (lecturer.professional_title or lecturer.position) if lecturer else ''
        identity.display_name = str(request.data.get('display_name') or identity.display_name or default_name)[:200]
        identity.position = str(request.data.get('position') or identity.position or default_position)[:220]
        identity.signature_image_path = rel
        identity.signature_image_sha256 = sha
        identity.is_active = True
        identity.save()
        identity = provision_identity(target)
        audit(request.user, 'update_signing_identity', 'user', target.id, image=name)
        return Response({'user_id': target.id, 'username': target.username, **identity_public_payload(identity)})
    except (DjangoValidationError, ValueError) as ex:
        return error_validation(ex) if isinstance(ex, DjangoValidationError) else failure(str(ex))


@api_view(['POST'])
def sign_approval(request, pk):
    guard = require_role(request, 'signer', 'admin')
    if guard:
        return guard
    password = str(request.data.get('password') or '')
    if not password or not request.user.check_password(password):
        return failure('Mật khẩu xác nhận không đúng.', 403)
    with transaction.atomic():
        a = ProfileApproval.objects.select_for_update().select_related('lecturer').filter(pk=pk).first()
        if not a:
            return failure('Không thấy hồ sơ.', 404)
        if a.status != ProfileApproval.Status.PRINT:
            return failure('Hồ sơ chưa ở trạng thái chờ ký số.', 409)
        identity = SigningIdentity.objects.filter(user=request.user, is_active=True).first()
        if not identity or not identity.encrypted_private_key or not identity.certificate_pem:
            return failure('Tài khoản chưa được cấp danh tính ký số nội bộ.')
        if not identity.signature_image_path:
            return failure('Chưa tải ảnh chữ ký để hiển thị trên BM-04.')
        signed_at = timezone.localtime()
        verification_code = f'SIG-U-{signed_at:%Y%m%d}-{secrets.token_hex(6).upper()}'
        display_name = identity.display_name or request.user.username
        sig_meta = {
            'image_path': str(read_private(identity.signature_image_path)),
            'image_rel_path': identity.signature_image_path,
            'signer_name': display_name,
            'signer_position': identity.position,
            'verification_code': verification_code,
            'signed_at': signed_at.strftime('%d/%m/%Y %H:%M:%S'),
        }
        personal_ds = a.digital_signatures.filter(signature_role=DigitalSignature.SignatureRole.PERSONAL).order_by('signed_at').first()
        personal_meta = dict(personal_ds.visible_signature or {}) if personal_ds else None
        if personal_meta and personal_meta.get('image_rel_path'):
            personal_meta['image_path'] = str(read_private(personal_meta['image_rel_path']))
        signed_pdf = profile_pdf(a.snapshot, a.code, draft=False, personal_signature=personal_meta, unit_signature=sig_meta)
        signature_b64 = sign_bytes(identity, signed_pdf)
        rel, pdf_sha = store_bytes(signed_pdf, 'signed_pdfs')
        ds = DigitalSignature.objects.create(
            approval=a, signer=request.user, identity=identity,
            signature_role=DigitalSignature.SignatureRole.UNIT,
            visible_signature=sig_meta,
            signature_b64=signature_b64, certificate_pem=identity.certificate_pem,
            source_pdf_sha256=a.pdf_sha256, signed_pdf_sha256=pdf_sha,
            signed_file_path=rel, verification_code=verification_code,
            verification_status='valid',
        )
        a.signed_pdf_path = rel
        a.signed_pdf_sha256 = pdf_sha
        a.signer = request.user
        a.status = ProfileApproval.Status.COMPLETED
        a.completed_at = timezone.now()
        a.save(update_fields=['signed_pdf_path','signed_pdf_sha256','signer','status','completed_at'])
        ApprovalEvent.objects.create(approval=a, actor=request.user, action='digital_signed', note=verification_code)
        audit(request.user, 'digital_sign_approval', 'approval', a.pk, verification_code=verification_code, signed_pdf_sha256=pdf_sha)
    return Response({**approval_out(a), 'verification_code': ds.verification_code})


@api_view(['GET'])
def verify_approval_signature(request, pk):
    a = ProfileApproval.objects.select_related('lecturer').filter(pk=pk).first()
    if not a:
        return failure('Không thấy hồ sơ.', 404)
    if has_role(request.user, 'lecturer') and a.lecturer.user_id != request.user.id:
        return failure('Không có quyền.', 403)
    signatures = list(a.digital_signatures.select_related('signer', 'identity').order_by('signed_at'))
    if not signatures:
        return failure('Hồ sơ chưa có chữ ký số nội bộ.', 404)
    results = []
    overall_valid = True
    for ds in signatures:
        path = read_private(ds.signed_file_path)
        file_exists = path.is_file()
        data = path.read_bytes() if file_exists else b''
        hash_ok = bool(file_exists and sha256_bytes(data) == ds.signed_pdf_sha256)
        crypto_ok = verify_bytes(ds.certificate_pem, data, ds.signature_b64) if hash_ok else False
        valid = bool(hash_ok and crypto_ok)
        overall_valid = overall_valid and valid
        if ds.verification_status != ('valid' if valid else 'invalid'):
            ds.verification_status = 'valid' if valid else 'invalid'
            ds.save(update_fields=['verification_status'])
        results.append({
            'valid': valid, 'signature_role': ds.signature_role,
            'signature_role_label': ds.get_signature_role_display(),
            'verification_code': ds.verification_code,
            'signer': ds.identity.display_name or ds.signer.username,
            'position': ds.identity.position,
            'certificate_serial': ds.identity.certificate_serial,
            'signed_at': ds.signed_at.isoformat(),
            'signed_pdf_sha256': ds.signed_pdf_sha256,
            'algorithm': ds.algorithm,
        })
    first = results[0]
    return Response({
        'valid': overall_valid,
        'verification_code': first['verification_code'],
        'signer': first['signer'],
        'position': first['position'],
        'certificate_serial': first['certificate_serial'],
        'signed_at': first['signed_at'],
        'signed_pdf_sha256': first['signed_pdf_sha256'],
        'algorithm': first['algorithm'],
        'signatures': results,
    })

def _report_range(request):
    period = request.query_params.get('period', 'month') if request else 'month'
    anchor_raw = request.query_params.get('date') if request else None
    try:
        anchor = date.fromisoformat(anchor_raw) if anchor_raw else timezone.localdate()
    except ValueError:
        anchor = timezone.localdate()

    if period == 'day':
        start_date, end_date = anchor, anchor + timedelta(days=1)
    elif period == 'week':
        start_date = anchor - timedelta(days=anchor.weekday())
        end_date = start_date + timedelta(days=7)
    elif period == 'year':
        start_date = date(anchor.year, 1, 1)
        end_date = date(anchor.year + 1, 1, 1)
    elif period == 'custom':
        try:
            start_date = date.fromisoformat(request.query_params.get('from', ''))
            end_inclusive = date.fromisoformat(request.query_params.get('to', ''))
            if end_inclusive < start_date:
                raise ValueError
            end_date = end_inclusive + timedelta(days=1)
        except ValueError:
            raise ValueError('Khoảng thời gian tùy chọn không hợp lệ.')
    else:
        period = 'month'
        start_date = date(anchor.year, anchor.month, 1)
        end_date = date(anchor.year + (1 if anchor.month == 12 else 0), 1 if anchor.month == 12 else anchor.month + 1, 1)

    tz = timezone.get_current_timezone()
    start_dt = timezone.make_aware(datetime.combine(start_date, datetime.min.time()), tz)
    end_dt = timezone.make_aware(datetime.combine(end_date, datetime.min.time()), tz)
    return period, start_date, end_date - timedelta(days=1), start_dt, end_dt


def report_data(request=None):
    try:
        period, start_date, end_date, start_dt, end_dt = _report_range(request)
    except ValueError as ex:
        return {'error': str(ex)}

    # V8: công trình không còn quy trình xác minh độc lập. Một công trình được
    # tính là "đã xác minh trong kỳ" khi mục WORK của nó được NCKH xác minh
    # trong checklist của ít nhất một hồ sơ BM-04.
    verified_work_ids = list(VerificationRecord.objects.filter(
        section=EvidenceDocument.Section.WORK,
        status=VerificationRecord.Status.VERIFIED,
        reviewed_at__gte=start_dt, reviewed_at__lt=end_dt,
    ).values_list('entity_id', flat=True).distinct())
    verified_period = ResearchWork.objects.filter(pk__in=verified_work_ids)
    by_year = list(verified_period.values('publication_year').annotate(total=Count('id')).order_by('-publication_year'))
    by_type = list(verified_period.values('work_type__name').annotate(total=Count('id')).order_by('-total', 'work_type__name'))
    by_department = list(
        WorkAuthor.objects.filter(
            work_id__in=verified_work_ids,
            lecturer__isnull=False, department_at_publication__isnull=False,
        ).values('department_at_publication__name').annotate(total=Count('work_id', distinct=True)).order_by('-total')
    )
    status_rows = list(ProfileApproval.objects.values('status').annotate(total=Count('id')).order_by('status'))
    return {
        'period': period,
        'from_date': str(start_date), 'to_date': str(end_date),
        'total_lecturers': Lecturer.objects.count(),
        'total_works': ResearchWork.objects.count(),
        'pending_works': ResearchWork.objects.filter(evidence__isnull=True).distinct().count(),
        'pending_approvals': ProfileApproval.objects.filter(status__in=['review', 'signer', 'print', 'signed']).count(),
        'approvals_created': ProfileApproval.objects.filter(submitted_at__gte=start_dt, submitted_at__lt=end_dt).count(),
        'approvals_approved': ProfileApproval.objects.filter(approved_at__gte=start_dt, approved_at__lt=end_dt).count(),
        'approvals_signed': DigitalSignature.objects.filter(signed_at__gte=start_dt, signed_at__lt=end_dt).count(),
        'approvals_completed': ProfileApproval.objects.filter(completed_at__gte=start_dt, completed_at__lt=end_dt).count(),
        'works_created': ResearchWork.objects.filter(created_at__gte=start_dt, created_at__lt=end_dt).count(),
        'works_verified': verified_period.count(),
        'active_users': AuditLog.objects.filter(at__gte=start_dt, at__lt=end_dt, actor__isnull=False).values('actor_id').distinct().count(),
        'approval_status': status_rows,
        'by_year': by_year, 'by_type': by_type, 'by_department': by_department,
        'note': 'V8 không xác minh công trình độc lập. Số công trình xác minh được tính từ checklist BM-04 của Phòng NCKH; công trình theo khoa dùng khoa tại thời điểm công bố.',
    }


@api_view(['GET'])
def reports(request):
    guard = require_role(request, 'reviewer', 'signer', 'admin', 'viewer_admin')
    if guard:
        return guard
    data = report_data(request)
    return failure(data['error']) if data.get('error') else Response(data)


@api_view(['GET'])
def reports_export(request):
    guard = require_role(request, 'reviewer', 'signer', 'admin', 'viewer_admin')
    if guard:
        return guard
    from openpyxl import Workbook
    from io import BytesIO
    r = report_data(request)
    if r.get('error'):
        return failure(r['error'])
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tong hop'
    for row in [['Chi tieu', 'So luong'], ['Tu ngay', r['from_date']], ['Den ngay', r['to_date']], ['Giang vien', r['total_lecturers']], ['Cong trinh duy nhat', r['total_works']], ['Cong trinh tao trong ky', r['works_created']], ['Cong trinh xac minh trong ky', r['works_verified']], ['BM-04 tao trong ky', r['approvals_created']], ['BM-04 da duyet trong ky', r['approvals_approved']], ['BM-04 da ky so trong ky', r['approvals_signed']], ['BM-04 hoan tat trong ky', r['approvals_completed']], ['Nguoi dung co hoat dong', r['active_users']], ['Cong trinh chua co minh chung', r['pending_works']], ['Ho so BM-04 dang xu ly', r['pending_approvals']]]:
        ws.append(row)
    for key, name, col in [('by_year', 'Theo nam', 'publication_year'), ('by_type', 'Theo loai', 'work_type__name'), ('by_department', 'Theo khoa', 'department_at_publication__name')]:
        s = wb.create_sheet(name)
        s.append(['Nhom', 'So cong trinh'])
        for item in r[key]:
            s.append([item[col], item['total']])
    wb.create_sheet('Ghi chu').append([r['note']])
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return FileResponse(out, as_attachment=True, filename='bao_cao_khoa_hoc.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _credentials_excel_base64(credentials):
    """Tạo file Excel mật khẩu tạm trong bộ nhớ, không lưu plaintext trên server."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tai_khoan_tam'
    ws.append(['staff_code', 'username', 'temporary_password', 'note'])
    for item in credentials:
        ws.append([
            item.get('staff_code', ''), item.get('username', ''),
            item.get('temporary_password', ''),
            'Yêu cầu đổi mật khẩu ngay lần đăng nhập đầu tiên',
        ])
    for col, width in {'A': 18, 'B': 24, 'C': 30, 'D': 48}.items():
        ws.column_dimensions[col].width = width
    out = BytesIO()
    wb.save(out)
    return base64.b64encode(out.getvalue()).decode('ascii')


@api_view(['GET', 'POST'])
def admin_users(request):
    guard = require_role(request, 'admin', 'viewer_admin') if request.method == 'GET' else require_role(request, 'admin')
    if guard:
        return guard
    if request.method == 'GET':
        return Response([{
            'id': u.pk,
            'username': u.username,
            'role': u.role,
            'is_active': u.is_active,
            'lecturer': u.lecturer.full_name if hasattr(u, 'lecturer') else '',
            'must_change_password': u.must_change_password,
            'last_login': u.last_login.isoformat() if u.last_login else None,
            'activation_status': 'pending' if u.must_change_password else 'active',
            'activation_status_label': 'Chưa kích hoạt/Chưa đổi mật khẩu' if u.must_change_password else 'Đã kích hoạt',
        } for u in User.objects.select_related('lecturer').order_by('username')[:1000]])
    d = json_body(request)
    try:
        if d.get('role') not in User.Role.values:
            return failure('Vai trò không hợp lệ.')
        if d.get('role') == User.Role.LECTURER:
            return failure('Tài khoản giảng viên phải được tạo qua nhập danh sách nhân sự để luôn gắn với hồ sơ Lecturer.')
        if not d.get('username'):
            return failure('Thiếu tên đăng nhập.')
        password = secrets.token_urlsafe(15)
        with transaction.atomic():
            u = User.objects.create_user(username=str(d['username']).strip(), password=password, role=d['role'], must_change_password=True)
            audit(request.user, 'create_user', 'user', u.pk)
    except IntegrityError:
        return failure('Tên tài khoản đã tồn tại.', 409)
    return Response({'id': u.id, 'username': u.username, 'temporary_password': password, 'warning': 'Mật khẩu chỉ hiển thị một lần; chuyển qua kênh bảo mật.'}, status=201)


@api_view(['POST'])
def admin_update_user(request, pk):
    guard = require_role(request, 'admin')
    if guard:
        return guard
    u = User.objects.filter(pk=pk).first()
    if not u:
        return failure('Không thấy người dùng.', 404)
    d = json_body(request)
    if 'is_active' in d:
        if u.pk == request.user.pk and not d['is_active']:
            return failure('Không thể tự khóa tài khoản hiện tại.')
        u.is_active = bool(d['is_active'])
    if 'role' in d and d['role'] in User.Role.values:
        u.role = d['role']
    u.save()
    audit(request.user, 'update_user', 'user', u.pk)
    return Response({'ok': True})


@api_view(['POST'])
def admin_reset_temporary_password(request, pk):
    guard = require_role(request, 'admin')
    if guard:
        return guard
    u = User.objects.filter(pk=pk).first()
    if not u:
        return failure('Không thấy người dùng.', 404)
    if not u.must_change_password:
        return failure('Tài khoản đã kích hoạt và đã đổi mật khẩu; chức năng này chỉ dùng cho tài khoản đang ở trạng thái chưa đổi mật khẩu.', 409)
    password = secrets.token_urlsafe(15)
    u.set_password(password)
    u.must_change_password = True
    u.save(update_fields=['password', 'must_change_password'])
    audit(request.user, 'reset_temporary_password', 'user', u.pk)
    return Response({
        'username': u.username,
        'temporary_password': password,
        'warning': 'Mật khẩu tạm chỉ hiển thị lần này và phải được đổi sau khi đăng nhập.',
    })


@api_view(['POST'])
def admin_import(request):
    guard = require_role(request, 'admin')
    if guard:
        return guard
    f = request.FILES.get('file')
    if not f or not f.name.lower().endswith('.xlsx'):
        return failure('Cần chọn tệp Excel .xlsx.')
    if f.size > settings_limit():
        return failure('Tệp Excel quá lớn.')
    from openpyxl import load_workbook
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        rows = wb.active.iter_rows(values_only=True)
        headers = [str(v or '').strip().lower() for v in next(rows)]
        required = ('staff_code', 'full_name', 'username', 'department_code')
        if any(k not in headers for k in required):
            return failure('Tệp Excel cần các cột: ' + ', '.join(required))
        parsed, problems, seen_codes, seen_users = [], [], set(), set()
        for i, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            r = dict(zip(headers, row))
            code = str(r.get('staff_code') or '').strip()
            name = str(r.get('full_name') or '').strip()
            username = str(r.get('username') or '').strip()
            dep = str(r.get('department_code') or '').strip()
            if not code or not name or not username or not dep:
                problems.append(f'Dòng {i}: thiếu trường bắt buộc')
                continue
            if code in seen_codes or username in seen_users:
                problems.append(f'Dòng {i}: trùng trong tệp')
                continue
            if not Department.objects.filter(code=dep).exists():
                problems.append(f'Dòng {i}: khoa {dep} chưa tồn tại')
                continue
            personal_id = str(r.get('personal_id') or '').strip()
            if personal_id and (personal_id in {str(x.get('personal_id') or '').strip() for x in parsed} or Lecturer.objects.filter(personal_id=personal_id).exists()):
                problems.append(f'Dòng {i}: số định danh cá nhân đã tồn tại')
                continue
            if Lecturer.objects.filter(staff_code=code).exists() or User.objects.filter(username=username).exists():
                problems.append(f'Dòng {i}: mã hoặc tài khoản đã tồn tại')
                continue
            parsed.append(r)
            seen_codes.add(code)
            seen_users.add(username)
            if len(parsed) > 1000:
                return failure('Chỉ nhập tối đa 1.000 tài khoản mỗi lần.')
    except Exception as ex:
        return failure('Không đọc được tệp Excel: ' + str(ex)[:120])
    if problems:
        return Response({'error': 'Tệp Excel có dữ liệu giảng viên không hợp lệ.', 'created': 0, 'errors': problems[:100], 'valid_rows': len(parsed)}, status=400)
    if request.data.get('dry_run', '0') == '1':
        return Response({'valid_rows': len(parsed), 'errors': [], 'preview': True})
    credentials = []
    with transaction.atomic():
        for r in parsed:
            password = secrets.token_urlsafe(15)
            u = User.objects.create_user(username=str(r['username']).strip(), password=password, role='lecturer', must_change_password=True)
            dob = None
            if r.get('date_of_birth'):
                if hasattr(r['date_of_birth'], 'date'):
                    dob = r['date_of_birth'].date()
                else:
                    dob = date_value(r['date_of_birth'])
            gender = str(r.get('gender') or '').strip().lower()
            if gender not in Lecturer.Gender.values:
                gender = ''
            department = Department.objects.get(code=str(r['department_code']).strip())
            lecturer = Lecturer.objects.create(
                user=u, staff_code=str(r['staff_code']).strip(), full_name=str(r['full_name']).strip(),
                personal_id=str(r.get('personal_id') or '').strip()[:30], date_of_birth=dob, gender=gender,
                professional_title=str(r.get('professional_title') or '').strip()[:220],
                position=str(r.get('position') or '').strip()[:220], home_address=str(r.get('home_address') or '').strip()[:700],
                phone=str(r.get('phone') or '').strip()[:40], email=str(r.get('email') or '').strip()[:254], department=department,
            )
            ScienceProfile.objects.create(lecturer=lecturer)
            DepartmentHistory.objects.create(
                lecturer=lecturer, department=department,
                start_date=date_value(r.get('department_start_date')) or timezone.localdate(),
            )
            credentials.append({'staff_code': lecturer.staff_code, 'username': u.username, 'temporary_password': password})
        audit(request.user, 'import_users', 'user', details_count=len(parsed))
    return Response({
        'created': len(credentials),
        'credentials': credentials,
        'credentials_xlsx_base64': _credentials_excel_base64(credentials),
        'credentials_filename': f'tai_khoan_tam_giang_vien_{timezone.localdate():%Y%m%d}.xlsx',
        'warning': 'Mật khẩu tạm chỉ được trả về ngay sau khi tạo và không được lưu dạng rõ. Hãy tải file Excel và chuyển qua kênh bảo mật.',
    })


def settings_limit():
    from django.conf import settings
    return settings.MAX_UPLOAD_BYTES


@api_view(['GET'])
def admin_departments_template(request):
    guard = require_role(request, 'admin', 'viewer_admin')
    if guard:
        return guard
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Don_vi'
    ws.append(['code', 'name', 'address', 'head_name', 'phone'])
    ws.append(['K01', 'Khoa mẫu', 'Địa chỉ đơn vị', 'Nguyễn Văn A', '0123456789'])
    for col, width in {'A': 18, 'B': 36, 'C': 42, 'D': 28, 'E': 20}.items():
        ws.column_dimensions[col].width = width
    out = BytesIO(); wb.save(out); out.seek(0)
    return FileResponse(out, as_attachment=True, filename='template_import_co_quan_phong_ban.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@api_view(['POST'])
def admin_departments_import(request):
    guard = require_role(request, 'admin')
    if guard:
        return guard
    f = request.FILES.get('file')
    if not f or not f.name.lower().endswith('.xlsx'):
        return failure('Cần chọn tệp Excel .xlsx.')
    if f.size > settings_limit():
        return failure('Tệp Excel quá lớn.')
    from openpyxl import load_workbook
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        rows = wb.active.iter_rows(values_only=True)
        headers = [str(v or '').strip().lower() for v in next(rows)]
        required = ('code', 'name')
        if any(k not in headers for k in required):
            return failure('Tệp Excel cần các cột: code, name, address, head_name, phone')
        parsed, problems, seen = [], [], set()
        for i, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            r = dict(zip(headers, row))
            code = str(r.get('code') or '').strip()
            name = str(r.get('name') or '').strip()
            if not code or not name:
                problems.append(f'Dòng {i}: thiếu code hoặc name')
                continue
            if code in seen:
                problems.append(f'Dòng {i}: trùng mã đơn vị trong tệp')
                continue
            conflict = Department.objects.filter(name=name).exclude(code=code).first()
            if conflict:
                problems.append(f'Dòng {i}: tên đơn vị đã được dùng bởi mã {conflict.code}')
                continue
            parsed.append({
                'code': code[:30], 'name': name[:180],
                'address': str(r.get('address') or '').strip()[:500],
                'head_name': str(r.get('head_name') or '').strip()[:200],
                'phone': str(r.get('phone') or '').strip()[:40],
            })
            seen.add(code)
            if len(parsed) > 1000:
                return failure('Chỉ nhập tối đa 1.000 đơn vị mỗi lần.')
    except Exception as ex:
        return failure('Không đọc được tệp Excel: ' + str(ex)[:120])
    if problems:
        return Response({'error': 'Tệp Excel cơ quan/phòng ban có dữ liệu không hợp lệ.', 'created': 0, 'updated': 0, 'valid_rows': len(parsed), 'errors': problems[:100]}, status=400)
    existing = {x.code for x in Department.objects.filter(code__in=[r['code'] for r in parsed])}
    if request.data.get('dry_run', '0') == '1':
        return Response({'valid_rows': len(parsed), 'created': len(parsed)-len(existing), 'updated': len(existing), 'errors': [], 'preview': True})
    created = updated = 0
    try:
        with transaction.atomic():
            for r in parsed:
                dep = Department.objects.filter(code=r['code']).first()
                if dep:
                    dep.name=r['name']; dep.address=r['address']; dep.head_name=r['head_name']; dep.phone=r['phone']; dep.save()
                    updated += 1
                else:
                    Department.objects.create(**r); created += 1
            audit(request.user, 'import_departments', 'department', details_count=len(parsed), created=created, updated=updated)
    except IntegrityError:
        return failure('Có mã hoặc tên đơn vị bị trùng trong cơ sở dữ liệu.', 409)
    return Response({'created': created, 'updated': updated, 'total': created+updated})


@api_view(['GET', 'POST'])
def admin_departments(request):
    guard = require_role(request, 'admin', 'viewer_admin') if request.method == 'GET' else require_role(request, 'admin')
    if guard:
        return guard
    if request.method == 'GET':
        return Response(list(Department.objects.values('id', 'code', 'name', 'address', 'head_name', 'phone').order_by('name')))
    d = json_body(request)
    if not d.get('code') or not d.get('name'):
        return failure('Thiếu mã hoặc tên khoa.')
    try:
        dep = Department.objects.create(
            code=str(d['code'])[:30], name=str(d['name'])[:180], address=str(d.get('address', ''))[:500],
            head_name=str(d.get('head_name', ''))[:200], phone=str(d.get('phone', ''))[:40],
        )
    except IntegrityError:
        return failure('Mã hoặc tên đã tồn tại.', 409)
    audit(request.user, 'create_department', 'department', dep.id)
    return Response({'id': dep.id, 'code': dep.code, 'name': dep.name}, status=201)


@api_view(['PATCH'])
def admin_department_detail(request, pk):
    guard = require_role(request, 'admin')
    if guard:
        return guard
    dep = Department.objects.filter(pk=pk).first()
    if not dep:
        return failure('Không thấy đơn vị.', 404)
    d = json_body(request)
    for field, limit in [('name', 180), ('code', 30), ('address', 500), ('head_name', 200), ('phone', 40)]:
        if field in d:
            setattr(dep, field, str(d[field] or '')[:limit])
    try:
        dep.save()
    except IntegrityError:
        return failure('Mã hoặc tên đơn vị bị trùng.', 409)
    audit(request.user, 'update_department', 'department', dep.id)
    return Response({'ok': True})


def _sync_employment_admin(l, records):
    if records is None:
        return
    if not isinstance(records, list) or len(records) > 200:
        raise ValueError('Danh sách quá trình công tác không hợp lệ.')
    kept=[]
    for i,x in enumerate(records):
        org=str(x.get('organization','')).strip()
        if not org: continue
        obj=l.employment_records.filter(pk=x.get('id')).first() if x.get('id') else None
        obj=obj or EmploymentRecord(lecturer=l)
        obj.start_year=year_value(x.get('start_year'), required=True)
        obj.end_year=year_value(x.get('end_year'))
        obj.position=str(x.get('position',''))[:220]
        obj.organization=org[:300]
        obj.organization_address=str(x.get('organization_address',''))[:500]
        obj.sort_order=i
        obj.save(); kept.append(obj.pk)
    if kept: l.employment_records.exclude(pk__in=kept).delete()
    else: l.employment_records.all().delete()


@api_view(['GET', 'PATCH'])
def admin_lecturer(request, pk):
    guard = require_role(request, 'admin', 'viewer_admin') if request.method == 'GET' else require_role(request, 'admin')
    if guard:
        return guard
    l = Lecturer.objects.select_related('department', 'user').filter(pk=pk).first()
    if not l:
        return failure('Không thấy giảng viên.', 404)
    if request.method == 'PATCH':
        d = json_body(request)
        before = {**lecturer_out(l), 'employment': [{
            'id': x.pk, 'start_year': x.start_year, 'end_year': x.end_year,
            'position': x.position, 'organization': x.organization, 'organization_address': x.organization_address,
        } for x in l.employment_records.all()]}
        try:
            if 'personal_id' in d:
                personal_id = str(d.get('personal_id') or '').strip()
                if personal_id and Lecturer.objects.exclude(pk=l.pk).filter(personal_id=personal_id).exists():
                    return failure('Số định danh cá nhân đã được sử dụng cho giảng viên khác.', 409)
            for field, limit in [('full_name', 200), ('personal_id', 30), ('professional_title', 220), ('position', 220), ('home_address', 700), ('phone', 40), ('email', 254)]:
                if field in d:
                    setattr(l, field, str(d[field] or '')[:limit])
            if 'gender' in d:
                if d['gender'] not in ('', *Lecturer.Gender.values):
                    return failure('Giới tính không hợp lệ.')
                l.gender = d['gender']
            if 'status' in d and d['status'] in Lecturer.Status.values:
                l.status = d['status']
            if 'date_of_birth' in d:
                l.date_of_birth = date_value(d['date_of_birth'])
            new_department_id = int(d.get('department_id', l.department_id))
            if new_department_id != l.department_id:
                dep = Department.objects.filter(pk=new_department_id).first()
                if not dep:
                    return failure('Khoa không hợp lệ.')
                start = date_value(d.get('effective_date'), required=True)
                if start > timezone.localdate():
                    return failure('Ngày có hiệu lực không được ở tương lai.')
                with transaction.atomic():
                    previous = l.department_history.filter(end_date__isnull=True).order_by('-start_date').first()
                    if previous and start <= previous.start_date:
                        return failure('Ngày chuyển khoa phải sau ngày bắt đầu giai đoạn hiện tại.')
                    if previous:
                        previous.end_date = start
                        previous.save(update_fields=['end_date'])
                    DepartmentHistory.objects.create(lecturer=l, department=dep, start_date=start)
                    l.department = dep
                    _sync_employment_admin(l, d.get('employment'))
                    l.save()
            else:
                with transaction.atomic():
                    _sync_employment_admin(l, d.get('employment'))
                    l.save()
            after = {**lecturer_out(l), 'employment': [{'id': x.pk, 'start_year': x.start_year, 'end_year': x.end_year, 'position': x.position, 'organization': x.organization, 'organization_address': x.organization_address} for x in l.employment_records.all()]}
            audit(request.user, 'update_lecturer', 'lecturer', l.pk, lecturer=l, before=before, after=after)
        except (ValueError, TypeError) as ex:
            return failure(str(ex))
    return Response({**lecturer_out(l), 'employment': [{
        'id': x.pk, 'start_year': x.start_year, 'end_year': x.end_year,
        'position': x.position, 'organization': x.organization, 'organization_address': x.organization_address,
    } for x in l.employment_records.all()]})


@api_view(['GET'])
def admin_lecturers(request):
    guard = require_role(request, 'admin', 'viewer_admin')
    if guard:
        return guard
    return Response([{**lecturer_out(l), 'employment': [{
        'id': x.pk, 'start_year': x.start_year, 'end_year': x.end_year,
        'position': x.position, 'organization': x.organization, 'organization_address': x.organization_address,
    } for x in l.employment_records.all()]} for l in Lecturer.objects.select_related('department', 'user').prefetch_related('employment_records').order_by('full_name')[:1000]])


@api_view(['GET'])
def audit_logs(request):
    guard = require_role(request, 'admin', 'viewer_admin', 'reviewer')
    if guard:
        return guard
    qs = AuditLog.objects.select_related('actor', 'lecturer').all()
    lecturer_id = request.query_params.get('lecturer_id')
    actor = str(request.query_params.get('actor') or '').strip()
    action = str(request.query_params.get('action') or '').strip()
    start_date = request.query_params.get('from')
    end_date = request.query_params.get('to')
    if lecturer_id: qs = qs.filter(lecturer_id=lecturer_id)
    if actor: qs = qs.filter(actor__username__icontains=actor)
    if action: qs = qs.filter(action__icontains=action)
    if start_date:
        try: qs = qs.filter(at__date__gte=date_value(start_date, required=True))
        except ValueError as ex: return failure(str(ex))
    if end_date:
        try: qs = qs.filter(at__date__lte=date_value(end_date, required=True))
        except ValueError as ex: return failure(str(ex))
    try:
        limit=min(max(int(request.query_params.get('limit','300')),1),1000)
    except ValueError:
        return failure('Giới hạn log không hợp lệ.')
    return Response([{
        'at': x.at.isoformat(), 'actor': x.actor.username if x.actor else None,
        'lecturer_id': str(x.lecturer_id) if x.lecturer_id else None,
        'lecturer': x.lecturer.full_name if x.lecturer else None,
        'action': x.action, 'entity': x.entity, 'entity_id': x.entity_id, 'details': x.details,
    } for x in qs[:limit]])


@api_view(['GET'])
def notifications(request):
    qs = Notification.objects.filter(recipient=request.user)[:100]
    return Response([{'id': n.id, 'title': n.title, 'body': n.body, 'created_at': n.created_at.isoformat(), 'read': bool(n.read_at)} for n in qs])


@api_view(['POST'])
def read_notifications(request):
    d = json_body(request)
    qs = Notification.objects.filter(recipient=request.user, read_at__isnull=True)
    if d.get('id'):
        qs = qs.filter(pk=d['id'])
    count = qs.update(read_at=timezone.now())
    return Response({'marked': count})
